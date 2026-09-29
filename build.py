"""cv.yaml -> ATS uyumlu PDF + DOCX (+ ATS'nin göreceği düz metin dökümü).

Kullanım:  python build.py
Çıktılar:  cikti/<dosya>.pdf, cikti/<dosya>.docx, cikti/<dosya>_ats_metin.txt
"""
import base64
import html
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "cv.yaml"
CSS = ROOT / "style.css"
OUT = ROOT / "cikti"

TR_ASCII = str.maketrans("çğıöşüÇĞİÖŞÜâÂ", "cgiosuCGIOSUaA")


# ---------------------------------------------------------------- cv.yaml -> iç yapı

class FormHatasi(Exception):
    pass


def txt(v):
    """YAML değerini temiz metne çevirir (sayılar, boş değerler, çok satırlı metin)."""
    if v is None:
        return ""
    return " ".join(str(v).split())


def need(item, key, where):
    val = txt(item.get(key)) if isinstance(item, dict) else ""
    if not val:
        raise FormHatasi(f"{where}: '{key}' alanı boş ya da eksik.")
    return val


def as_list(v):
    if v is None:
        return []
    return [txt(x) for x in (v if isinstance(v, list) else [v]) if txt(x)]


def dates(item):
    return " – ".join(x for x in (txt(item.get("baslangic")), txt(item.get("bitis"))) if x)


def entries(items, where, left, sub=None, extra=None):
    """Liste bölümlerini (iş, eğitim, proje) başlık/tarih/alt satır/madde kayıtlarına çevirir."""
    blocks = []
    for i, it in enumerate(items or [], 1):
        w = f"{where} #{i}"
        if not isinstance(it, dict):
            raise FormHatasi(f"{w}: alanlar 'anahtar: değer' şeklinde olmalı.")
        bullets = (extra(it) if extra else []) + as_list(it.get("maddeler"))
        sub_text = sub(it) if sub else ""
        blocks.append({"type": "entry", "left": left(it, w), "right": dates(it),
                       "sub": (sub_text, "") if sub_text else None, "bullets": bullets})
    return blocks


def load_cv(data):
    if not isinstance(data, dict):
        raise FormHatasi("cv.yaml boş ya da hatalı.")
    g = data.get("genel_bilgiler") or {}
    cv = {"name": need(g, "ad_soyad", "genel_bilgiler"), "photo": txt(g.get("fotograf")) or None,
          "header": [], "sections": []}

    contact = [txt(g.get(k)) for k in ("sehir", "telefon", "email")]
    links = [f"{label}: {txt(g.get(k))}" for k, label in
             (("github", "GitHub"), ("linkedin", "LinkedIn"), ("web", "Web")) if txt(g.get(k))]
    for line in (txt(g.get("unvan")), " | ".join(c for c in contact if c), " | ".join(links)):
        if line:
            cv["header"].append(line)

    def section(title, blocks):
        if blocks:
            cv["sections"].append({"title": title, "blocks": blocks})

    def bullet_list(items):
        items = [x for x in items if x]
        return [{"type": "list", "items": items}] if items else []

    if txt(data.get("ozet")):
        section("Profesyonel Özet", [{"type": "para", "text": txt(data["ozet"])}])

    section("İş Deneyimi", entries(
        data.get("is_deneyimi"), "is_deneyimi",
        left=lambda it, w: need(it, "pozisyon", w),
        sub=lambda it: txt(it.get("kurum"))))

    section("Eğitim", entries(
        data.get("egitim"), "egitim",
        left=lambda it, w: ", ".join(x for x in (txt(it.get("derece")), need(it, "bolum", w)) if x),
        sub=lambda it: txt(it.get("okul")),
        extra=lambda it: [f"Genel Not Ortalaması: {txt(it['ortalama'])}"] if txt(it.get("ortalama")) else []))

    pubs = []
    for i, it in enumerate(data.get("yayinlar") or [], 1):
        parts = [f"**{need(it, 'baslik', f'yayinlar #{i}').rstrip('.')}.**"]
        if txt(it.get("yazarlar")):
            parts.append(txt(it["yazarlar"]).rstrip(".") + ".")
        place = ", ".join(x for x in (txt(it.get("yayin_yeri")), txt(it.get("yil"))) if x)
        if place:
            parts.append(place + ".")
        if txt(it.get("link")):
            parts.append(txt(it["link"]))
        if txt(it.get("kod")):
            parts.append(f"Kod: {txt(it['kod'])}")
        pubs.append(" ".join(parts))
    section("Yayınlar", bullet_list(pubs))

    section("Projeler", entries(
        data.get("projeler"), "projeler",
        left=lambda it, w: need(it, "isim", w),
        sub=lambda it: ", ".join(as_list(it.get("linkler")))))

    skills = data.get("beceriler") or {}
    if not isinstance(skills, dict):
        raise FormHatasi("beceriler: 'Kategori: [a, b, c]' şeklinde olmalı.")
    section("Beceriler", bullet_list(
        f"**{txt(k)}:** {', '.join(as_list(v))}" for k, v in skills.items() if as_list(v)))

    langs = [f"**{need(it, 'ad', 'diller_ve_sinavlar')}:** {txt(it.get('sonuc'))}".rstrip()
             for it in data.get("diller_ve_sinavlar") or []]
    if langs:
        section("Yabancı Dil ve Sınavlar", [{"type": "para", "text": " | ".join(langs)}])

    certs = []
    for i, it in enumerate(data.get("sertifikalar") or [], 1):
        c = need(it, "ad", f"sertifikalar #{i}")
        if txt(it.get("kurum")):
            c += f" – {txt(it['kurum'])}"
        if txt(it.get("yil")):
            c += f" ({txt(it['yil'])})"
        certs.append(c)
    section("Sertifikalar", bullet_list(certs))

    refs = []
    for i, it in enumerate(data.get("referanslar") or [], 1):
        role = ", ".join(x for x in (txt(it.get("unvan")), txt(it.get("kurum"))) if x)
        refs.append(" – ".join(x for x in (f"**{need(it, 'ad', f'referanslar #{i}')}**", role,
                                           txt(it.get("email")), txt(it.get("telefon"))) if x))
    section("Referanslar", bullet_list(refs))
    return cv


def read_cv():
    import yaml
    try:
        data = yaml.safe_load(SRC.read_text(encoding="utf-8"))
    except yaml.YAMLError as e:
        mark = getattr(e, "problem_mark", None)
        where = f" (satır {mark.line + 1}, sütun {mark.column + 1})" if mark else ""
        raise FormHatasi(f"cv.yaml okunamadı{where}: {getattr(e, 'problem', e)}. "
                         "Girintileri ve ': ' içeren değerlerin tırnakta olduğunu kontrol et.")
    return load_cv(data)


INLINE = re.compile(r"\*\*(.+?)\*\*|\[(.+?)\]\((.+?)\)")
# Düz yazılmış e-posta ve adresler (github.com/kullanici, https://...) otomatik link olur.
AUTOLINK = re.compile(
    r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"
    r"|https?://[^\s|,;()]+"
    r"|(?:www\.)?[a-z0-9-]+(?:\.[a-z0-9-]+)*\.(?:com|io|dev|org|net|me|ai|app|tr)(?:/[^\s|,;()]*)?",
    re.I)


def autolink_url(t):
    if "@" in t and "/" not in t:
        return "mailto:" + t
    return t if re.match(r"https?://", t, re.I) else "https://" + t


def autolink(text):
    parts, pos = [], 0
    for m in AUTOLINK.finditer(text):
        if m.start() > pos:
            parts.append(("text", text[pos:m.start()], None))
        parts.append(("link", m.group(0), autolink_url(m.group(0))))
        pos = m.end()
    if pos < len(text):
        parts.append(("text", text[pos:], None))
    return parts


def inline_parts(text):
    """Metni (tür, metin, url) parçalarına böler: tür 'text' | 'bold' | 'link'."""
    parts, pos = [], 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            parts.extend(autolink(text[pos:m.start()]))
        if m.group(1) is not None:
            parts.append(("bold", m.group(1), None))
        else:
            parts.append(("link", m.group(2), m.group(3)))
        pos = m.end()
    if pos < len(text):
        parts.extend(autolink(text[pos:]))
    return parts


def photo_bytes(cv):
    """Fotoğrafı küçültüp JPEG olarak döndürür (PDF/DOCX boyutu şişmesin)."""
    if not cv["photo"]:
        return None
    path = (ROOT / cv["photo"])
    if not path.exists():
        print(f"! Fotoğraf bulunamadı: {path}")
        return None
    try:
        from PIL import Image
        im = Image.open(path).convert("RGB")
        im.thumbnail((500, 600))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=88, optimize=True)
        return buf.getvalue()
    except ImportError:
        return path.read_bytes()


# ---------------------------------------------------------------- HTML / PDF

def inline_html(text):
    out = []
    for kind, t, url in inline_parts(text):
        t = html.escape(t)
        if kind == "bold":
            out.append(f"<strong>{t}</strong>")
        elif kind == "link":
            out.append(f'<a href="{html.escape(url)}">{t}</a>')
        else:
            out.append(t)
    return "".join(out)


def to_html(cv, photo=None):
    h = [f"<header><div class=\"head-text\"><h1>{html.escape(cv['name'])}</h1>"]
    for i, line in enumerate(cv["header"]):
        cls = "title" if i == 0 else "contact"
        h.append(f'<p class="{cls}">{inline_html(line)}</p>')
    h.append("</div>")
    if photo:
        b64 = base64.b64encode(photo).decode()
        h.append(f'<img class="photo" alt="" src="data:image/jpeg;base64,{b64}">')
    h.append("</header>")

    for sec in cv["sections"]:
        h.append(f"<section><h2>{html.escape(sec['title'])}</h2>")
        for b in sec["blocks"]:
            if b["type"] == "para":
                h.append(f"<p>{inline_html(b['text'])}</p>")
            elif b["type"] == "list":
                h.append("<ul>" + "".join(f"<li>{inline_html(i)}</li>" for i in b["items"]) + "</ul>")
            else:
                h.append('<div class="entry"><div class="row">'
                         f'<h3>{inline_html(b["left"])}</h3>'
                         f'<span class="date">{inline_html(b["right"])}</span></div>')
                if b["sub"]:
                    sl, sr = b["sub"]
                    h.append(f'<div class="row sub"><span>{inline_html(sl)}</span>'
                             f'<span>{inline_html(sr)}</span></div>')
                if b["bullets"]:
                    h.append("<ul>" + "".join(f"<li>{inline_html(i)}</li>" for i in b["bullets"]) + "</ul>")
                h.append("</div>")
        h.append("</section>")

    css = CSS.read_text(encoding="utf-8")
    return ("<!doctype html><html lang=\"tr\"><head><meta charset=\"utf-8\">"
            f"<title>{html.escape(cv['name'])} - CV</title><style>{css}</style></head>"
            f"<body>{''.join(h)}</body></html>")


def find_browser():
    candidates = [
        os.environ.get("CV_BROWSER", ""),
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    ]
    for c in candidates:
        if c and Path(c).exists():
            return c
    for name in ("google-chrome", "chromium", "chromium-browser", "chrome", "msedge"):
        p = shutil.which(name)
        if p:
            return p
    return None


def build_pdf(html_path, pdf_path):
    browser = find_browser()
    if not browser:
        print("! Chrome/Edge bulunamadı, PDF atlandı (CV_BROWSER ortam değişkeniyle yol verebilirsin).")
        return False
    with tempfile.TemporaryDirectory() as profile:
        cmd = [browser, "--headless=new", "--disable-gpu", "--no-first-run",
               f"--user-data-dir={profile}", "--no-pdf-header-footer",
               "--print-to-pdf-no-header", f"--print-to-pdf={pdf_path}",
               html_path.as_uri()]
        subprocess.run(cmd, check=False, capture_output=True, timeout=120)
    return pdf_path.exists()


# ---------------------------------------------------------------- DOCX

def build_docx(cv, docx_path, photo=None):
    try:
        from docx import Document
        from docx.enum.text import WD_TAB_ALIGNMENT
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement, parse_xml
        from docx.oxml.ns import nsdecls, qn
        from docx.shared import Mm, Pt
    except ImportError:
        print("! python-docx yok, DOCX atlandı (pip install python-docx).")
        return False

    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Mm(297), Mm(210)
        s.top_margin = s.bottom_margin = Mm(10)
        s.left_margin = s.right_margin = Mm(14)
    text_width = doc.sections[0].page_width - Mm(28)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(11.6)

    for name, size in (("Title", 22), ("Heading 1", 12.5), ("Heading 2", 10)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Calibri", Pt(size), True
        st.font.color.rgb = None
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), "Calibri")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        st.font.italic = False
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(6)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(2)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(3)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(0)
    doc.styles["Title"].paragraph_format.space_after = Pt(2)
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for bdr in title_ppr.findall(qn("w:pBdr")):  # varsayılan mavi alt çizgiyi kaldır
        title_ppr.remove(bdr)
    bullet_fmt = doc.styles["List Bullet"].paragraph_format
    bullet_fmt.left_indent, bullet_fmt.first_line_indent = Mm(5), Mm(-3.5)

    def add_hyperlink(p, text, url, bold=False, italic=False):
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        r, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
        for on, tag in ((bold, "w:b"), (italic, "w:i")):
            if on:
                rpr.append(OxmlElement(tag))
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(rpr)
        r.append(t)
        h.append(r)
        p._p.append(h)

    def add_runs(p, text, bold=False, italic=False):
        for kind, t, url in inline_parts(text):
            if kind == "link":
                add_hyperlink(p, t, url, bold=bold, italic=italic)
            else:
                r = p.add_run(t)
                r.bold = True if (bold or kind == "bold") else None
                r.italic = italic or None

    def float_photo(p, data, height):
        """Fotoğrafı sağ üst köşeye sabitler; metin solundan akar (tablo kullanmadan)."""
        p.add_run().add_picture(io.BytesIO(data), height=height)
        inline = p._p.xpath(".//wp:inline")[0]
        ext, pr = inline.find(qn("wp:extent")), inline.find(qn("wp:docPr"))
        anchor = parse_xml(
            f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="114300" distR="0" simplePos="0" '
            'relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/>'
            '<wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH>'
            '<wp:positionV relativeFrom="margin"><wp:posOffset>0</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{ext.get("cx")}" cy="{ext.get("cy")}"/>'
            '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="left"/>'
            f'<wp:docPr id="{pr.get("id")}" name="Fotograf"/><wp:cNvGraphicFramePr/></wp:anchor>')
        anchor.append(inline.find(qn("a:graphic")))
        inline.getparent().replace(inline, anchor)

    def bottom_border(p):
        ppr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
            b.set(qn(k), v)
        bdr.append(b)
        ppr.append(bdr)

    def lr_para(left, right, style=None, italic=False):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        add_runs(p, left, italic=italic)
        if right:
            p.add_run("\t")
            r = p.add_run(right)
            r.bold = False
            r.italic = italic or None
        return p

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        add_runs(p, text)

    p = doc.add_paragraph(style="Title")
    p.add_run(cv["name"])
    if photo:
        float_photo(p, photo, Mm(29))
    for i, line in enumerate(cv["header"]):
        p = doc.add_paragraph()
        add_runs(p, line)
        if i == 0:
            for r in p.runs:
                r.font.size = Pt(12)

    for sec in cv["sections"]:
        bottom_border(doc.add_paragraph(sec["title"], style="Heading 1"))
        for b in sec["blocks"]:
            if b["type"] == "para":
                p = doc.add_paragraph()
                add_runs(p, b["text"])
            elif b["type"] == "list":
                for it in b["items"]:
                    bullet(it)
            else:
                p = lr_para(b["left"], b["right"], style="Heading 2")
                p.paragraph_format.keep_with_next = True
                if b["sub"]:
                    sp = lr_para(*b["sub"], italic=True)
                    sp.paragraph_format.keep_with_next = True
                for it in b["bullets"]:
                    bullet(it)

    doc.core_properties.title = f"{cv['name']} - CV"
    doc.core_properties.author = cv["name"]
    doc.save(docx_path)
    return True


# ---------------------------------------------------------------- ATS kontrolü

def dump_pdf_text(pdf_path, txt_path):
    try:
        import pymupdf
    except ImportError:
        return None
    with pymupdf.open(pdf_path) as d:
        text = "\n".join(page.get_text() for page in d)
        pages = d.page_count
    txt_path.write_text(text, encoding="utf-8")
    return pages


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    try:
        cv = read_cv()
    except FormHatasi as e:
        print(f"! HATA: {e}")
        return 1
    OUT.mkdir(exist_ok=True)
    base = "cv_" + re.sub(r"[^a-z0-9]+", "_", cv["name"].translate(TR_ASCII).lower()).strip("_")

    html_path = OUT / f"{base}.html"
    pdf_path = OUT / f"{base}.pdf"
    docx_path = OUT / f"{base}.docx"
    txt_path = OUT / f"{base}_ats_metin.txt"

    photo = photo_bytes(cv)
    html_path.write_text(to_html(cv, photo), encoding="utf-8")
    if pdf_path.exists():
        pdf_path.unlink()

    ok_pdf = build_pdf(html_path, pdf_path)
    ok_docx = build_docx(cv, docx_path, photo)

    print(f"HTML : {html_path}")
    if ok_pdf:
        pages = dump_pdf_text(pdf_path, txt_path)
        print(f"PDF  : {pdf_path}" + (f"  ({pages} sayfa)" if pages else ""))
        if pages:
            print(f"ATS  : {txt_path}  (ATS'nin PDF'ten okuyacağı düz metin)")
    if ok_docx:
        print(f"DOCX : {docx_path}")
    return 0 if ok_pdf or ok_docx else 1


if __name__ == "__main__":
    sys.exit(main())
